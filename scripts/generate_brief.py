#!/usr/bin/env python3
"""Generate costForensics platform brief as PPTX presentation."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# --- Colors ---
DARK_BLUE = RGBColor(0x1A, 0x23, 0x7E)
ORANGE = RGBColor(0xFF, 0x6F, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
DARK_TEXT = RGBColor(0x21, 0x21, 0x21)
SUBTLE_TEXT = RGBColor(0x75, 0x75, 0x75)

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=DARK_TEXT, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=DARK_TEXT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_after = Pt(6)
        p.level = 0
    return tf


def add_footer(slide, text="costForensics — Commerce Financial Forensics Platform"):
    add_textbox(slide, Inches(0.5), Inches(6.9), Inches(12), Inches(0.4),
                text, font_size=10, color=SUBTLE_TEXT, alignment=PP_ALIGN.LEFT)


def add_accent_bar(slide):
    shape = slide.shapes.add_shape(
        1, Inches(0.5), Inches(1.45), Inches(1.5), Inches(0.06)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ORANGE
    shape.line.fill.background()


def make_slide(prs, title_text, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, WHITE)

    add_textbox(slide, Inches(0.5), Inches(0.5), Inches(12), Inches(0.9),
                title_text, font_size=36, color=DARK_BLUE, bold=True)
    add_accent_bar(slide)

    if subtitle:
        add_textbox(slide, Inches(0.5), Inches(1.7), Inches(12), Inches(0.6),
                    subtitle, font_size=20, color=SUBTLE_TEXT)
        bullet_top = Inches(2.5)
    else:
        bullet_top = Inches(2.0)

    add_bullet_list(slide, Inches(0.8), bullet_top, Inches(11), Inches(4.5),
                    bullets, font_size=18, color=DARK_TEXT)
    add_footer(slide)
    return slide


def make_two_column_slide(prs, title_text, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    add_textbox(slide, Inches(0.5), Inches(0.5), Inches(12), Inches(0.9),
                title_text, font_size=36, color=DARK_BLUE, bold=True)
    add_accent_bar(slide)

    # Left column
    add_textbox(slide, Inches(0.8), Inches(1.9), Inches(5.5), Inches(0.5),
                left_title, font_size=22, color=ORANGE, bold=True)
    add_bullet_list(slide, Inches(1.0), Inches(2.5), Inches(5.3), Inches(4),
                    left_items, font_size=16, color=DARK_TEXT)

    # Right column
    add_textbox(slide, Inches(7.0), Inches(1.9), Inches(5.5), Inches(0.5),
                right_title, font_size=22, color=ORANGE, bold=True)
    add_bullet_list(slide, Inches(7.2), Inches(2.5), Inches(5.3), Inches(4),
                    right_items, font_size=16, color=DARK_TEXT)

    add_footer(slide)
    return slide


def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # ── Slide 1: Cover ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BLUE)
    add_textbox(slide, Inches(1), Inches(2.0), Inches(11), Inches(1.2),
                "costForensics", font_size=54, color=WHITE, bold=True,
                alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(3.3), Inches(11), Inches(0.8),
                "Commerce Financial Forensics Platform", font_size=28,
                color=ORANGE, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1), Inches(4.5), Inches(11), Inches(0.6),
                "Unified cloud costs \u2022 commerce operations \u2022 margin forensics",
                font_size=18, color=WHITE, alignment=PP_ALIGN.CENTER)

    # ── Slide 2: Problem ──
    make_slide(prs, "The Problem", [
        "\u2022  Cloud costs scattered across AWS, GCP, and Azure consoles",
        "\u2022  Commerce margins calculated in spreadsheets with no audit trail",
        "\u2022  Silent margin erosion: untracked discounts, commission drift, hidden fees",
        "\u2022  No single source of truth connecting cloud spend to product profitability",
        "\u2022  Finance teams discover problems weeks or months after they occur",
    ], subtitle="No unified visibility into cloud costs and commerce margins")

    # ── Slide 3: Solution ──
    make_slide(prs, "The Solution", [
        "\u2022  One platform: cloud costs + commerce + financial ledger + forensics",
        "\u2022  Real-time margin decomposition from revenue to net profit",
        "\u2022  SHA-256 verified forensic audit trail for every financial event",
        "\u2022  Automated anomaly detection across cloud spend and margin drift",
        "\u2022  Multi-tenant SaaS with database-per-tenant isolation",
    ], subtitle="costForensics: see what your margins actually are")

    # ── Slide 4: Cloud Cost Analysis ──
    make_two_column_slide(prs, "Cloud Cost Analysis",
        "Cost Tracking", [
            "\u2022  AWS, GCP, Azure account management",
            "\u2022  Per-service cost records (EC2, S3, Lambda, etc.)",
            "\u2022  Multi-currency support with exchange rates",
            "\u2022  Automated CSV import and cloud API sync",
        ],
        "Controls", [
            "\u2022  Budgets with threshold alerts (50%, 80%, 100%)",
            "\u2022  Anomaly detection (spike, drift, trend, outlier)",
            "\u2022  Severity classification (low, medium, high, critical)",
            "\u2022  Monthly and quarterly cost reports",
        ])

    # ── Slide 5: Commerce Operations ──
    make_two_column_slide(prs, "Commerce Operations",
        "Catalog & Parties", [
            "\u2022  Products with SKU, EAN, UPC identifiers",
            "\u2022  Cost and price tracking per product",
            "\u2022  Sellers with configurable commission rates",
            "\u2022  Customer segmentation (enterprise, SMB, consumer)",
        ],
        "Order Lifecycle", [
            "\u2022  Full order management: pending \u2192 shipped \u2192 delivered",
            "\u2022  Multi-item orders with quantity and pricing",
            "\u2022  Cancellation and refund workflows",
            "\u2022  Per-order cost, revenue, and margin tracking",
        ])

    # ── Slide 6: Financial Engine ──
    make_two_column_slide(prs, "Financial Engine",
        "Ledger & Audit", [
            "\u2022  Double-entry ledger (revenue, COGS, discounts, fees)",
            "\u2022  SHA-256 hash chain for tamper-evident records",
            "\u2022  Forensic events: price changes, discount applications",
            "\u2022  Full traceability from event to ledger entry",
        ],
        "Promotions & Payments", [
            "\u2022  Promotion rules: percentage, flat, tiered, buy-X-get-Y",
            "\u2022  Funding split: seller-funded, platform-funded, shared",
            "\u2022  Bidirectional payments (inbound + outbound)",
            "\u2022  Multi-currency exchange rate management",
        ])

    # ── Slide 7: Margin Forensics ──
    make_slide(prs, "Margin Forensics", [
        "\u2022  Per-order decomposition: Revenue \u2192 COGS \u2192 Discounts \u2192 Commission \u2192 Fees \u2192 Net Margin",
        "\u2022  Gross margin and net margin percentage calculation",
        "\u2022  Margin drift detection: alerts when margins deviate from baselines",
        "\u2022  Drill-down by product, seller, customer segment, and time period",
        "\u2022  Forensic trail links every margin component to its source event",
    ], subtitle="See exactly where every cent of margin goes")

    # ── Slide 8: Analytics ──
    make_two_column_slide(prs, "Analytics & Reporting",
        "P&L Analysis", [
            "\u2022  Drill-down by day, month, quarter, year",
            "\u2022  Revenue, COGS, gross profit, operating expenses",
            "\u2022  Net income with trend visualization",
            "\u2022  Comparative period analysis",
        ],
        "Operational Intelligence", [
            "\u2022  Customer retention cohort analysis",
            "\u2022  Margin drift reports with severity tracking",
            "\u2022  Cloud cost anomaly trending",
            "\u2022  Budget utilization dashboards",
        ])

    # ── Slide 9: Data Integration ──
    make_slide(prs, "Data Integration", [
        "\u2022  CSV import with validation, error reporting, and progress tracking",
        "\u2022  Shopify connector: product and order sync",
        "\u2022  WooCommerce connector: catalog and transaction import",
        "\u2022  MedusaJS connector: headless commerce integration",
        "\u2022  Webhook support for real-time event ingestion",
        "\u2022  Batch sync for historical data backfill",
    ], subtitle="Connect your commerce stack in minutes")

    # ── Slide 10: Architecture ──
    make_two_column_slide(prs, "Architecture",
        "Design Principles", [
            "\u2022  Hexagonal / clean architecture",
            "\u2022  Rich domain models with enforced invariants",
            "\u2022  Database-per-tenant isolation",
            "\u2022  Separate GORM models with mapper layer",
        ],
        "Tech Stack", [
            "\u2022  Backend: Go, Gin, GORM, PostgreSQL",
            "\u2022  Frontend: React, TypeScript, Vite",
            "\u2022  UI: Material UI, Recharts",
            "\u2022  State: Zustand + TanStack Query",
        ])

    # ── Slide 11: API at a Glance ──
    make_slide(prs, "API at a Glance", [
        "\u2022  70+ REST endpoints across 6 domains",
        "\u2022  JWT authentication with tenant-scoped authorization",
        "\u2022  Domains: Cloud Costs, Commerce, Ledger, Forensics, Analytics, Integration",
        "\u2022  Middleware chain: Recovery \u2192 RequestID \u2192 Logging \u2192 CORS \u2192 Auth \u2192 TenantDB",
        "\u2022  OpenAPI/Swagger documentation",
    ], subtitle="Production-ready API design")

    # ── Slide 12: Summary ──
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK_BLUE)
    add_textbox(slide, Inches(1), Inches(0.8), Inches(11), Inches(1),
                "Why costForensics?", font_size=42, color=WHITE, bold=True,
                alignment=PP_ALIGN.CENTER)

    differentiators = [
        "\u2022  Forensic audit trail — SHA-256 hash chain, tamper-evident, fully traceable",
        "\u2022  Margin reality engine — see actual margins, not estimates",
        "\u2022  Unified view — cloud costs and commerce in one platform",
        "\u2022  Multi-tenant SaaS — database-per-tenant isolation, zero cross-tenant risk",
        "\u2022  Open architecture — clean hexagonal design, easy to extend and integrate",
    ]
    add_bullet_list(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(3.5),
                    differentiators, font_size=20, color=WHITE)

    add_textbox(slide, Inches(1), Inches(5.8), Inches(11), Inches(0.8),
                "costForensics — Know your real margins.",
                font_size=24, color=ORANGE, bold=True, alignment=PP_ALIGN.CENTER)

    return prs


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    output_dir = os.path.join(project_root, "docs")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "costForensics_brief.pptx")

    prs = build_presentation()
    prs.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    main()
